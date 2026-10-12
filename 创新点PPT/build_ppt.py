"""Build an editable innovation deck using reviewed images and recorded results."""
from pathlib import Path
import hashlib
import json
import sys
sys.path.insert(0, '/tmp/avs_ppt_deps')
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

ROOT=Path(__file__).resolve().parent.parent
OUT=ROOT/'创新点PPT'
SCENE=OUT/'素材/新版/scene'; VISION=OUT/'素材/新版/vision'; CONTROL=OUT/'素材/新版/control'
RUN_NAME='speed_mixed_020_20261012_run2'; RUN=ROOT/'avs_ws/acceptance'/RUN_NAME
BG='101B2C'; CARD='1D2C42'; WHITE='F4F7FC'; MUTED='A6B7CE'; TEAL='42D5BB'; GOLD='FFC466'; BLUE='7AABFF'
r=Presentation(); r.slide_width=Inches(13.333333); r.slide_height=Inches(7.5)
r.core_properties.title='AVS作品特色创新点'; r.core_properties.author='AVS项目'
slides=[]
metadata={}
for folder in (SCENE,VISION,CONTROL):
    data=json.loads((folder/'manifest.json').read_text())
    for asset in (data['assets'] if isinstance(data,dict) else data):
        path=Path(asset['file'])
        full=ROOT/path if path.parts[:1]==('创新点PPT',) else folder/path
        metadata[str(full)]=asset

def box(s,x,y,w,h,color=CARD,kind=MSO_SHAPE.ROUNDED_RECTANGLE):
    sh=s.shapes.add_shape(kind,Inches(x),Inches(y),Inches(w),Inches(h))
    sh.fill.solid();sh.fill.fore_color.rgb=RGBColor.from_string(color);sh.line.fill.background()
    if kind==MSO_SHAPE.ROUNDED_RECTANGLE: sh.adjustments[0]=.08
    return sh

def txt(s,text,x,y,w,h,size=20,color=WHITE,bold=False,align=PP_ALIGN.LEFT):
    sh=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=sh.text_frame
    tf.word_wrap=True;tf.margin_left=tf.margin_right=Inches(.02);tf.margin_top=tf.margin_bottom=0
    tf.vertical_anchor=MSO_ANCHOR.TOP
    for i,line in enumerate(text.split('\n')):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=line;p.alignment=align
        p.font.name='Noto Sans CJK SC';p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=RGBColor.from_string(color)
        p.space_after=Pt(5);p.line_spacing=1.1
    return sh

def new(title,sub,source,notes):
    s=r.slides.add_slide(r.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=RGBColor.from_string(BG)
    txt(s,'AVS  /  救援无人车 · 特色创新',.48,.2,10,.3,11,TEAL)
    txt(s,title,.48,.68,12.4,.65,29,bold=True);txt(s,sub,.48,1.4,12.4,.48,16,MUTED)
    box(s,.5,6.99,12.3,.012,'33465F',MSO_SHAPE.RECTANGLE)
    txt(s,source,.5,7.1,11.8,.23,8.6,MUTED);txt(s,f'{len(slides)+1:02d}',12.35,7.06,.42,.28,11,TEAL,True,PP_ALIGN.RIGHT)
    slides.append(dict(page=len(slides)+1,title=title,source=source,notes=notes,pictures=[]));return s

def pic(s,path,x,y,w,h,meaning,kind=None,source=None):
    path=Path(path)
    with Image.open(path) as im: iw,ih=im.size
    scale=min(w/iw,h/ih);nw,nh=iw*scale,ih*scale;left,top=x+(w-nw)/2,y+(h-nh)/2
    s.shapes.add_picture(str(path),Inches(left),Inches(top),Inches(nw),Inches(nh))
    meta=metadata.get(str(path),{})
    slides[-1]['pictures'].append(dict(file=str(path.relative_to(ROOT)),meaning=meaning,
        kind=kind or meta.get('kind',meta.get('type','实际图片')),source=source or meta.get('sources',meta.get('source',str(path.relative_to(ROOT)))),
        sha256=hashlib.sha256(path.read_bytes()).hexdigest(),width=iw,height=ih,asset_metadata=meta))
    return left,top,nw,nh

def flow(s,items,x=.5,y=6.0,w=12.3,h=.64,size=17):
    gap=.26;cell=(w-gap*(len(items)-1))/len(items)
    for i,t in enumerate(items):
        left=x+i*(cell+gap);box(s,left,y,cell,h);txt(s,t,left+.1,y+.13,cell-.2,h-.17,size,TEAL,True,PP_ALIGN.CENTER)
        if i<len(items)-1: box(s,left+cell+.035,y+h/2-.07,.17,.14,TEAL,MSO_SHAPE.CHEVRON)

def fact(s,title,value,x,y,w,h=1.,color=TEAL):
    box(s,x,y,w,h);txt(s,value,x+.16,y+.1,w-.32,.46,24,color,True);txt(s,title,x+.16,y+.64,w-.32,.28,13,MUTED)

def callout(s,text,x,y,w,px,py):
    txt(s,text,x,y,w,.36,16,TEAL,True)
    c=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x+w*.5),Inches(y+.4),Inches(px),Inches(py))
    c.line.color.rgb=RGBColor.from_string(TEAL);c.line.width=Pt(1.2)
    box(s,px-.03,py-.03,.06,.06,TEAL,MSO_SHAPE.OVAL)

s=new('六项创新，贯通观察、清障、定位与返航','面向狭窄救援场地的工程创新｜问题、机制与证据对应',
      '图：v19 原始连续录屏 90 s；内容核对当前源码及 2026-10-12 验收。',
      '六项贡献围绕实际救援流程组织。ROS 2、Gazebo、Nav2、YOLOv5、AprilTag和PnP属于采用的既有技术。场景图来自v19历史成功录像，符合当前world布局，不称最新提速截图。')
pic(s,SCENE/'scene_overview_actual.png',.5,2.02,6.05,4.55,'灰走廊、蓝起点、黄核心区、红箱的真实Gazebo全景')
txt(s,'真实仿真场地  /  3 × 3 m',.7,6.59,5.8,.28,13,MUTED)
topics=[('01','云台扫描与多帧确认','观察动作与任务闭环'),('02','推障机构与控制协同','被动推板、闭环与接触边界'),
        ('03','移动障碍地图残影处理','激光自由证据修正导航图'),('04','YOLOv5 与多级核验','图案、几何与时序共同判断'),
        ('05','任务阶段的速度精度协同','普通路段提速，精细阶段低速'),('06','时间同步定位与地图标注','拍摄时刻 TF 与稳健位置融合')]
for i,(num,title,sub) in enumerate(topics):
    x,y=6.85+(i%2)*3.,2.08+(i//2)*1.5;box(s,x,y,2.85,1.34)
    txt(s,num,x+.14,y+.11,.5,.32,16,TEAL,True);txt(s,title,x+.14,y+.51,2.58,.36,16,WHITE,True);txt(s,sub,x+.14,y+.97,2.58,.27,11,MUTED)

s=new('01  云台扫描与多帧确认协同','底盘沿路线行驶，云台负责侧墙观察；所需目标确认后任务继续',
      '图：最终 run2 / image_0380、0395；代码：camera_scan_node、hazard_tracker、mission_node。',
      '12秒固定正弦扫描由真实物理关节驱动，未实现按置信度自动选视角。两图属于同一停车点不同扫视时刻，位置差约2.9mm、偏航差约1.6°，不能作为连续三个处理帧证据。危险标识空间关联后连续三个处理帧才发布确认；缺帧重置计数。停车先检查至少0.8秒且实际速度足够低，确认后观察至少1秒；28秒超时停车。')
for x,file,title,frame in [( .5,'scan_stop_A.png','同一停车点 · 扫视时刻 A','0380'),(4.82,'scan_stop_B.png','同一停车点 · 扫视时刻 B','0395')]:
    pic(s,VISION/file,x,2.02,4.15,3.1,title);txt(s,title,x,5.17,4.15,.33,15,TEAL,True);txt(s,f'完整原图 / image_{frame}.png',x,5.58,4.15,.25,10.5,MUTED)
fact(s,'固定云台扫描周期','12 s',9.22,2.18,3.55);fact(s,'危险标识任务确认','连续 3 帧',9.22,3.42,3.55)
txt(s,'当前帧显示框\n→ 连续确认\n→ 发布任务事件',9.42,4.74,3.15,1.0,17)
flow(s,['到观察点','验证底盘已停','云台左右观察','收到确认后继续'])

s=new('02  整车结构与闭环推进协同','被动 V 形推板降低机构复杂度；机械引导与里程计反馈共同完成推进',
      '主图：当前 model.sdf 等比例渲染；小图：v19 原始录屏 175 s 的真实完整车辆。',
      '主图按当前SDF实际几何、位置、配色渲染，明确是模型示意，不称Gazebo截图。真实整车小图原始98×65像素，仅作来源参照。横梁宽0.28m、轮缘最大宽0.285m、门宽0.50m。斜翼对小偏差提供引导趋势，不是任意偏移自定心或力控制。里程计闭环限制横偏4cm、航向0.12rad；超限中止。')
box(s,.5,2.,8.1,4.73)
left,top,nw,nh=pic(s,SCENE/'vehicle_sdf_front.png',.72,2.15,7.67,4.2,'完整车辆、驱动轮、V形推板、雷达和云台相机','按实际 SDF 渲染示意')
callout(s,'云台相机',6.4,2.03,1.55,left+nw*.65,top+nh*.37);callout(s,'激光雷达',3.9,2.03,1.65,left+nw*.52,top+nh*.29)
callout(s,'驱动轮',.8,5.82,1.3,left+nw*.34,top+nh*.71);callout(s,'V 形推板',6.45,5.82,1.65,left+nw*.8,top+nh*.78)
txt(s,'按当前 SDF 等比例渲染 · 整车示意',1.32,6.4,6.4,.24,12,MUTED,align=PP_ALIGN.CENTER)
box(s,8.87,2.,3.93,1.68);pic(s,SCENE/'vehicle_actual_crop.png',9.1,2.16,1.55,1.04,'真实完整整车裁剪，原始98×65像素')
txt(s,'真实整车参照',10.84,2.25,1.78,.38,15,TEAL,True);txt(s,'原录屏尺寸有限',10.84,2.82,1.78,.35,11,MUTED)
fact(s,'推板宽 / 门洞宽','0.28 / 0.50 m',8.87,3.95,3.93);fact(s,'允许横偏 / 航向偏差','4 cm / 0.12 rad',8.87,5.25,3.93,color=GOLD)

s=new('02  真实接触推障：对准、推进、脱离','箱体位移来自 Gazebo 动力学接触；车辆推进距离由里程计闭环控制',
      '图：同一次 v19 原始连续录屏 128 / 135 / 160 s；参数读取当前 mission_params.yaml。',
      '三图仅展示v19同一次历史成功运行机械动作，不与最新提速数据混为同次。1.08m是车辆沿初始航向的里程计推进目标，不是箱体实际位移。推箱上限0.10m/s、24秒超时，检查侧偏、航向及推进距离。推完执行有界转向脱离，红箱不被瞬移或删除。')
pic(s,SCENE/'push_sequence_actual.png',.5,2.02,12.3,3.98,'同一次历史成功运行的推前、推中、推后真实场景')
fact(s,'推障速度上限','0.10 m/s',.5,6.,3.92,h=.89);fact(s,'车体里程计推进目标','1.08 m',4.69,6.,3.92,h=.89);fact(s,'推障超时即停车','24 s',8.88,6.,3.92,h=.89,color=GOLD)

s=new('02  有界接触与完整轮廓停车保护','推板、车轮一并检查；主动推障与普通避碰采用受限的接触边界',
      '图：当前 footprint 与 motion_safety.py 参数绘制的原理示意；非实测扫掠轨迹。',
      '真实凸多边形车体按反应时间和保守制动时长扫掠，图中未绘3.5cm外侧余量。安全检查位于速度平滑后、驱动前，指令与未裁剪实测速度分别检查。普通阶段完整保护；推障仅放行低速前方窄窗口；脱离例外固定在原箱体区域，最长12秒、平移6cm且近零线速。失效即停，不宣称连续安全证明或目标视觉身份验证。')
pic(s,CONTROL/'完整轮廓停车保护示意.png',.5,2.04,12.3,3.12,'真实footprint直行与转弯扫掠原理')
for i,(a,b) in enumerate([('普通导航','完整轮廓 + 双速度检查'),('授权推障','低速前方接触，侧后仍保护'),('转向脱离','例外固定：≤12 s / ≤6 cm')]):
    x=.5+i*4.19;box(s,x,5.38,3.93,1.31);txt(s,a,x+.16,5.54,3.61,.38,19,TEAL,True);txt(s,b,x+.16,6.13,3.61,.42,15)

s=new('03  激光自由证据驱动的地图残影处理','清除旧箱体留下的静态占用，保持结构墙和当前位置障碍的保护',
      '图：实际 rescue_map.pgm + navigation_map.yaml 允许区域；右侧为处理原理。',
      '左图为实际PGM依据YAML原点和分辨率显示，叠加两个配置区域。没有/map与/navigation_map真实同区域前后记录，因此仅解释原理，不称实测前后。限定区域内累计三次不同时间戳扫描的自由证据，射线停在量测障碍前8cm；原PGM不改，派生图用于导航，实时障碍层标记箱体新位置。这不是通用动态SLAM。')
pic(s,CONTROL/'实际地图与允许清理区域.png',.5,1.98,5.45,4.93,'实际保存地图叠加允许清理区域')
for i,(a,b) in enumerate([('旧占用残影','箱体移走，保存地图仍可能堵路'),('自由证据累积','同一栅格累计 3 次独立扫描'),('派生导航图','清理限定区域，实时层标记新障碍')]):
    y=2.12+i*1.42;box(s,6.3,y,6.45,1.12);txt(s,f'{i+1:02d}  {a}',6.51,y+.12,5.96,.38,21,TEAL,True);txt(s,b,6.51,y+.66,5.96,.34,16)
    if i<2:box(s,9.34,y+1.16,.17,.2,TEAL,MSO_SHAPE.DOWN_ARROW)
txt(s,'处理原理 · 不是实测栅格前后对照',6.5,6.57,6.1,.27,12,MUTED)

s=new('04  轻量 YOLOv5 与图案、几何、时序核验','同色负例训练 + 牌面验证 + 局部放大重检，减少场景干扰并恢复部分小目标',
      '图：同帧 v19 / image_0462 算法回放；结果记录 vision_recovery_replay.json。',
      '旧权重曾把黄色车身、蓝起点、红箱误报为辐射、电气、火灾。当前训练加入正确同色物体与空场景负例，候选核验白色图案、颜色和几何后拟合牌边，再空间关联与连续三帧确认。此图左半与v19 image_0462逐像素对应，右半为同输入回放；记录before=[]、after=[ELECTRICAL]。全图阈值0.45，已核验区域局部重检下限0.25。仍有截断/斜视瞬时漏框，不以此单例宣称全视频准确率。')
pic(s,VISION/'recovery_same_frame_verified.png',.5,2.,8.1,3.86,'ELECTRICAL同一输入帧的区域恢复回放','同帧离线算法回放')
txt(s,'同一输入帧：局部放大后重新进行 YOLO 分类',.66,5.94,7.78,.35,15,TEAL,True)
for i,(a,b) in enumerate([('训练端','小目标、斜视、模糊增强\n同色物体与空场景负例'),('判断端','颜色与白色图案核验\n拟合实际牌面边界'),('恢复端','已核验区域裁剪放大\n再次 YOLO 分类确认')]):
    y=2.04+i*1.42;box(s,8.85,y,3.94,1.28);txt(s,a,9.,y+.09,3.62,.32,18,TEAL,True);txt(s,b,9.,y+.53,3.62,.68,15)
flow(s,['YOLO 候选','图案与牌边核验','必要时区域重检','空间关联与确认'],y=6.37,h=.45,size=14)

s=new('04  实际识别画面：五类标识与真实红箱','五类均选择实际分类框可见的原图；紫色 PUSH BARRIER 来自独立红箱检测器',
      '图：最终 run2 / image_0092、0077、0213、0335、0320、0775；完整原图拼排。',
      '五类来自同一次最终成功运行且实际有绿色分类框。第六图0775紫框对应真实红箱，非YOLO FIRE分类。相机图像不重画框、不改坐标；仅展示样例，不证明逐帧零误报或连续确认。')
for i,(file,title,num) in enumerate([('fire.png','FIRE · 火灾标识',92),('chemical.png','CHEMICAL · 化学标识',77),('radiation.png','RADIATION · 辐射标识',213),('electrical.png','ELECTRICAL · 电气标识',335),('biohazard.png','BIOHAZARD · 生物危害',320),('real_red_box.png','真实红箱 · PUSH BARRIER',775)]):
    x,y=.5+(i%3)*4.19,1.98+(i//3)*2.48;box(s,x,y,3.93,2.35)
    txt(s,title,x+.12,y+.08,3.69,.33,16,TEAL,True);pic(s,VISION/file,x+.13,y+.48,3.67,1.65,title+'真实框原图')
    txt(s,f'run2 / image_{num:04d}.png',x+.13,y+2.13,3.67,.2,9.5,MUTED)

s=new('05  面向任务阶段的速度与精度协同','普通路段提高效率，窄门与精细操作保留低速；位置到达和朝向对正分阶段',
      '图：最终 run2 实际位置采样 + 当前阶段配置；颜色是阶段速度上限，不是实测速率。',
      '轨迹由run2独立验收车辆位置和当前world墙体绘制，控制不读真值。颜色表示阶段参数，实际会减速、停车和转向。普通0.20m/s、5cm到点容差、0.20m减速距离；精细0.12m/s、2.5cm、0.40m。推障0.10m/s。精细步骤包含推障对正、返航门洞接近/转向/穿门、最终起点。到位置后读当前TF坐标再对正，减少AMCL微调引发的追旧终点。按预设任务阶段切换，不称在线最优速度。')
pic(s,CONTROL/'实测轨迹与速度档位.png',.5,1.98,5.04,4.91,'同一次成功运行的实际轨迹，按任务阶段染色')
for i,(a,b,c) in enumerate([('普通路段','0.20 m/s','到点容差 5 cm · 减速距离 0.20 m'),('窄门 / 推障对正 / 最终起点','0.12 m/s','到点容差 2.5 cm · 减速距离 0.40 m'),('闭环推障','0.10 m/s','横偏、航向、距离与超时共同约束')]):
    y=2.07+i*1.25;box(s,5.9,y,6.87,1.11);txt(s,a,6.08,y+.12,4.05,.33,18,TEAL,True);txt(s,b,10.4,y+.11,2.15,.38,22,GOLD,True,PP_ALIGN.RIGHT);txt(s,c,6.08,y+.66,6.45,.3,15)
flow(s,['先到达位置','读取当前位置','再对正朝向'],x=5.9,y=6.05,w=6.87,h=.63,size=15)

s=new('05  完整任务提速 7.6%，保留识别确认条件','评价完整任务耗时：包含行驶、停车、观察、推障、对正与返航',
      '数据：speed_tuning_validation.json；1 次原速基准，2 次相同最终配置完整成功运行。',
      '首次导航至完成310.8198→287.1957秒，省23.6241秒即7.6006%；含启动317.3463→292.5701秒，不混采集总时长。两次均5类+4标签齐全，22视觉文件哈希与基准一致。图像视角不同，参数一致和任务成功不证明逐帧精确率/召回率完全不变。0.24分档候选296.84秒更慢，两个更高速度失败候选不计成绩。两次定位最高误差12.83和6.83cm，不宣称精度不变或全场景最优。')
pic(s,CONTROL/'完整任务耗时对照.png',.5,2.1,7.43,4.5,'同口径310.82→287.20秒，减少7.6%')
fact(s,'两次完整运行的任务确认','5 类 + 4 标签',8.22,2.12,4.56);fact(s,'视觉相关文件与基准一致','22 个文件哈希',8.22,3.39,4.56)
txt(s,'保留的观察条件',8.43,4.75,4.15,.37,20,TEAL,True);txt(s,'12 s 云台扫描\n3 帧危险标识确认\n至少 4 次有效标签观测\n停车稳定与确认后观察',8.43,5.24,4.15,1.44,16)

s=new('06  运动云台下的时间同步定位与地图标注','图像匹配拍摄时刻的实际云台姿态，避免旧图像与新姿态错配',
      '图：最终 run2 / image_0639、0745、0761、0870 的局部裁剪；代码：vision_node.py。',
      '图像非阻塞等待拍摄时间戳TF，过期丢弃，实际云台关节角经joint_states进TF。标签已知0.1875m边长及四角用IPPE Square PnP估计位置，重投影门限2.5px；同图时间戳转map。至少4次有效观测，最近20样本中值、离群拒绝及0.35指数平滑；缺帧不重置，不称连续四帧。标签灾情名由预设ID映射，不称独立识别真实毒气/人员险情。不读取标签模型真值。裁剪保留原算法ID、坐标和框，不重写。')
for i,(num,frame) in enumerate([(0,639),(1,745),(2,761),(3,870)]):
    x,y=.5+(i%2)*3.99,1.98+(i//2)*2.29;box(s,x,y,3.77,2.15);txt(s,f'T{num} · 地图坐标',x+.1,y+.08,3.54,.28,15,TEAL,True)
    pic(s,VISION/f'tag_{num}_crop.png',x+.12,y+.45,3.52,1.45,f'T{num}实际ID和地图坐标','原图局部裁剪')
    txt(s,f'image_{frame:04d}.png / 裁背景',x+.12,y+1.93,3.52,.18,9,MUTED)
for i,(a,b) in enumerate([('时间同步','拍摄时刻 → 实际关节姿态 TF'),('视觉测距','已知标签尺寸 + 四角 PnP'),('质量筛选','重投影门限 + 离群点拒绝'),('稳健地图位置','有效观测 + 中值与平滑')]):
    y=2.11+i*1.07;txt(s,a,8.85,y,3.94,.36,18,TEAL,True);txt(s,b,8.85,y+.48,3.94,.51,15)
txt(s,'图像 ID 与 RViz Marker 一一对应',8.85,6.49,3.94,.31,14,GOLD)

ev=json.loads((RUN/'evaluation.json').read_text());first=json.loads((ROOT/'avs_ws/acceptance/speed_mixed_020_20261012_run1/evaluation.json').read_text())
errs=list(ev['core_marker_errors_m'].values());e1=list(first['core_marker_errors_m'].values())
s=new('工程特色：独立、可追溯的全流程验收','任务、视觉与控制不读取世界真值；验收旁路记录实际结果与源码版本',
      '图与主指标：最终 run2；重复结果：最终 run1；165 项回归测试见 speed_tuning_validation。',
      '主图和主指标同属最终run2，蓝去程、绿回程、紫标签估计，非RViz截图。run2五类四标签全，箱体1.004902m、返航5.253cm、Marker4.409–6.826cm。run1相同配置通过，Marker7.671–12.830cm，诚实显示波动。两次93源文件哈希一致，22视觉文件与基准一致，165回归测试通过。车体与墙体离散采样无重叠不是连续无碰撞证明。当前只验证已知仿真场地和合成图案。')
pic(s,SCENE/'trajectory_latest.png',.5,2.,4.36,4.9,'最终run2实际轨迹、标签估计和箱体初终位置')
fact(s,'走廊类别 / 核心标签','5 / 5  ·  4 / 4',5.15,2.08,3.67);fact(s,'箱体实际位移',f'{ev["barrier_displacement_m"]:.4f} m',9.12,2.08,3.67)
fact(s,'返航位置误差',f'{ev["return_error_m"]*100:.2f} cm',5.15,3.35,3.67);fact(s,'四 Marker 平面误差',f'{min(errs)*100:.2f}–{max(errs)*100:.2f} cm',9.12,3.35,3.67)
box(s,5.15,4.67,7.64,1.14);txt(s,'同配置重复运行：两次完整验收均通过',5.34,4.8,7.24,.37,20,TEAL,True)
txt(s,f'第 1 次 Marker：{min(e1)*100:.2f}–{max(e1)*100:.2f} cm；第 2 次：{min(errs)*100:.2f}–{max(errs)*100:.2f} cm',5.34,5.32,7.24,.32,15)
txt(s,'图像 / 事件 / 轨迹 / 源码哈希组成证据链',5.34,6.,7.24,.39,19,WHITE,True)
txt(s,'165 项回归测试通过；采样无重叠不等于连续无碰撞证明。',5.34,6.51,7.24,.33,12,MUTED)

for slide,info in zip(r.slides,slides):
    image_lines='\n'.join(f'- {p["meaning"]}｜{p["kind"]}｜{p["source"]}' for p in info['pictures'])
    slide.notes_slide.notes_text_frame.text=info['notes']+'\n\n图片与证据来源：\n'+image_lines+'\n\n'+info['source']
    for shape in slide.shapes:
        if shape.left<0 or shape.top<0 or shape.left+shape.width>r.slide_width+100 or shape.top+shape.height>r.slide_height+100:
            raise ValueError(f'Shape outside slide {info["page"]}: {shape.shape_type}')
r.save(OUT/'AVS作品特色创新点.pptx')
(OUT/'PPT配图来源.json').write_text(json.dumps(dict(version='2026-10-12',innovations=6,slides=slides),ensure_ascii=False,indent=2))
print('Saved',len(slides),'slides:',OUT/'AVS作品特色创新点.pptx')
