from pathlib import Path
import sys
sys.path.insert(0, '/tmp/avs_ppt_deps')
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

ROOT=Path(__file__).resolve().parent.parent
OUT=ROOT/'创新点PPT'; AS=OUT/'素材'
E=ROOT/'avs_ws/acceptance/recording_yolov5_smac_v19/evidence'
FONT='/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc'
BG='101B2C'; CARD='1D2C42'; WHITE='F4F7FC'; MUTED='A6B7CE'; TEAL='42D5BB'; GOLD='FFC466'

# 实际验收帧，保持完整画面，不重画检测框。
cells=[('0113','FIRE｜红色警示图案'),('0161','CHEMICAL｜橙色图案'),('0322','RADIATION｜黄色图案'),('0528','ELECTRICAL｜蓝色图案'),('0588','BIOHAZARD｜紫色图案'),('0416','蓝色起点与墙面图案'),('0573','黄色车身与红色箱体'),('1607','红箱、蓝色起点与图案')]
canvas=Image.new('RGB',(1920,960),'#'+BG); d=ImageDraw.Draw(canvas)
for i,(num,cap) in enumerate(cells):
 x=24+(i%4)*474;y=18+(i//4)*470
 im=Image.open(E/f'image_{num}.png').convert('RGB'); im.thumbnail((450,338),Image.Resampling.LANCZOS)
 canvas.paste(im,(x,y+45));d.text((x,y),cap,font=ImageFont.truetype(FONT,22),fill='#'+WHITE)
 d.text((x,y+393),f'v19 / image_{num}.png',font=ImageFont.truetype(FONT,17),fill='#'+MUTED)
d.text((24,915),'实际验收画面：五类识别与同色干扰场景；不是旧权重前后对比或逐帧准确率统计。',font=ImageFont.truetype(FONT,23),fill='#'+MUTED)
canvas.save(AS/'YOLOv5识别与同色干扰拼图.png')

r=Presentation();r.slide_width=Inches(13.333);r.slide_height=Inches(7.5)
def box(s,x,y,w,h,color=CARD,kind=MSO_SHAPE.ROUNDED_RECTANGLE):
 sh=s.shapes.add_shape(kind,Inches(x),Inches(y),Inches(w),Inches(h));sh.fill.solid();sh.fill.fore_color.rgb=RGBColor.from_string(color);sh.line.fill.background();return sh
def txt(s,text,x,y,w,h,size=20,color=WHITE,bold=False):
 sh=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=sh.text_frame;tf.word_wrap=True
 for j,line in enumerate(text.split('\n')):
  p=tf.paragraphs[0] if j==0 else tf.add_paragraph();p.text=line;p.font.name='Noto Sans CJK SC';p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=RGBColor.from_string(color)
 return sh
def slide(title,sub,source,note):
 s=r.slides.add_slide(r.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=RGBColor.from_string(BG)
 txt(s,'AVS  /  救援无人车 · 特色创新',.45,.2,10,.3,12,TEAL)
 txt(s,title,.45,.7,12.4,.65,30,bold=True);txt(s,sub,.45,1.42,12.4,.5,16,MUTED)
 txt(s,source,.45,7.03,11.9,.3,10,MUTED);txt(s,f'{len(r.slides):02d}',12.4,7.02,.5,.3,12,TEAL)
 s.notes_slide.notes_text_frame.text=note+'\n\n证据来源：'+source
 return s
def picture(s,p,x,y,w,h):
 im=Image.open(p);iw,ih=im.size;scale=min(w/iw,h/ih);nw,nh=iw*scale,ih*scale
 s.shapes.add_picture(str(p),Inches(x+(w-nw)/2),Inches(y+(h-nh)/2),width=Inches(nw),height=Inches(nh))
def flow(s,items,y=5.6):
 w=(12.3-.25*(len(items)-1))/len(items)
 for i,t in enumerate(items):
  x=.5+i*(w+.25);box(s,x,y,w,.8);txt(s,t,x+.12,y+.13,w-.24,.55,17,TEAL,True)
  if i<len(items)-1:txt(s,'→',x+w,y+.2,.3,.4,17)

s=slide('四个创新点，贯通感知、动作与导航','面向狭窄救援场地的工程创新｜从发现目标到清障、入室与返航','来源：项目源码、学习指南第 7 章；场地图为已有仿真截图。','开场：我们的特色在于围绕任务中实际出现的问题，把感知、机构、控制和地图处理组织为完整流程。四点分别回答怎么看、怎么推、推完怎么走、怎么辨别。')
picture(s,ROOT/'avs_ws/docs/rescue_navigation_verified.png',.5,2,6.4,4.8)
for i,(a,b) in enumerate([('01  主动扫描','云台扫描 × 多帧确认'),('02  机构控制协同','被动 V 形推板 × 闭环推进'),('03  移动障碍地图处理','激光自由证据 × 残影清理'),('04  抗同色干扰识别','轻量 YOLOv5 × 图案与几何核验')]):
 box(s,7.2,2+i*1.16,5.6,1.02);txt(s,a,7.4,2.08+i*1.16,5.2,.35,19,TEAL,True);txt(s,b,7.4,2.48+i*1.16,5.2,.4,18)

s=slide('01  云台主动扫描与多帧识别结合','创新重点：把“车辆行驶”和“侧墙观察”分开，形成主动观察闭环','来源：camera_scan_node.py、vision_node.py、hazard_tracker.py；图片为 v19 验收帧。','底盘按路线走，物理云台承担左右观察。实际关节角进入 TF，图像使用曝光时刻变换。观察点停车等待识别，同一目标连续三个处理帧才确认，缺帧或空间跳变重新计数。这里强调观察动作与任务协同，第四点强调识别算法链。')
for i,(num,label) in enumerate([('0207','远处观察'),('0447','接近目标'),('0528','当前帧识别')]):
 picture(s,E/f'image_{num}.png',.5+i*4.2,2.1,4.0,3);txt(s,label,.6+i*4.2,5.02,3.8,.4,18,TEAL)
flow(s,['底盘沿路线行驶','云台左右扫描','观察点等待确认','连续 3 帧发布事件'])
txt(s,'实际云台角 → TF；曝光时间戳 → 图像与姿态对应',.6,6.55,12,.35,16,MUTED)

s=slide('02  推障机构与控制协同','创新重点：被动机械引导 + 里程计闭环 + 有界接触安全','来源：模型 SDF、mission_node.py、motion_safety.py；整车图按实际 model.sdf 几何与配色渲染。','固定 V 形推板减少执行器复杂度，斜翼对小偏差具有引导趋势。软件限制横偏与航向误差，接触授权限定在车头窄窗口、低速、时间和距离内。不是力控，也不宣称任意偏移自动自定心。v19 独立验收箱体实际移动 0.9749 米。')
picture(s,AS/'整车结构展示.png',.5,2.05,5.25,3.15)
txt(s,'整车结构：推板 · 轮组 · 雷达 · 云台',.65,5.15,5.2,.4,17,TEAL)
for i,(a,b) in enumerate([('结构匹配','推板 0.28 m / 门宽 0.50 m'),('闭环约束','横偏 ≤ 4 cm / 航向误差 ≤ 0.12 rad'),('接触约束','低速、有限时间与距离；侧后仍受保护')]):
 txt(s,a,6.15,2.12+i*1.08,6.5,.4,21,TEAL,True);txt(s,b,6.15,2.57+i*1.08,6.5,.55,18)
flow(s,['对准箱体','授权低速接触','闭环推进','脱离并恢复导航'])
txt(s,'同次 v19 仿真验收：箱体实际位移 0.9749 m',.6,6.55,12,.35,17,GOLD)

s=slide('03  针对移动障碍的地图残影处理','创新重点：限定区域内用激光自由证据修正导航图，避免旧箱体位置持续堵路','来源：navigation_map_node.py；栅格图为处理原理示意。','推走箱体以后，原始静态地图中的旧占用格仍可能挡住门口。方案只在场地先验限定区域中，累计三个不同扫描时间戳的自由证据后，将对应栅格置为空闲。原始地图保持不变，发布派生导航图，实时障碍层仍标记箱体新位置。不是通用动态 SLAM。')
for k,label in enumerate(['旧位置留下占用残影','激光累计自由证据','派生地图恢复通路']):
 x=.6+k*4.2;txt(s,label,x,2.04,3.9,.45,20,TEAL,True)
 for row in range(7):
  for col in range(9):
   c='263B53'
   if col in (0,8) or row in (0,6):c='718096'
   if row in (3,4) and col in (4,5):c='D84A55' if k==0 else ('FFC466' if k==1 else '42D5BB')
   box(s,x+col*.39,2.78+row*.36,.34,.31,c,MSO_SHAPE.RECTANGLE)
 txt(s,['旧箱体占用','累计 3 次空闲证据','清理旧位；新位仍避障'][k],x,5.55,3.9,.7,18)
txt(s,'限定区域  +  多次扫描证据  +  保留结构墙  +  实时代价层',.65,6.5,12,.4,19,GOLD)

s=slide('04  轻量 YOLOv5 与多级稳定确认','创新重点：从“颜色像目标”走向“图案、边界与时序共同可信”','来源：权重 hazard_yolov5.json；yolov5_detector.py、panel_geometry.py、hazard_tracker.py；历史问题见学习指南。','第四点可以作为场景化视觉工程创新。旧权重曾将黄色车身、蓝色起点和红箱误报为 RADIATION、ELECTRICAL、FIRE，并出现框超出图案或跳变。当前做法补入同色负例，并联合图案、面板几何和连续帧确认。没有同条件前后量化对照，不声称误报降低某个百分比。')
for i,(a,b) in enumerate([('历史问题','黄色车身 → RADIATION\n蓝色起点 → ELECTRICAL\n红色箱体 → FIRE\n候选框过大 / 在候选间跳变'),('针对性改进','同色物体与空场景负例\n轻量三尺度检测，320 输入\n图案核验 + 面板轮廓拟合\n空间关联 + 连续 3 帧确认'),('本组贡献','比赛图案训练与数据增强\n几何边界与时序验证链\n检测、显示和任务事件集成\n面向比赛场景的工程创新')]):
 x=.5+i*4.2;box(s,x,2.2,4,2.85);txt(s,a,x+.2,2.42,3.6,.5,23,TEAL,True);txt(s,b,x+.2,3.03,3.6,1.9,18)
flow(s,['YOLO 候选','图案核验','牌边拟合','跨帧关联','确认事件'])
txt(s,'YOLOv5 为现有架构；特色在于训练与验证链的场景化改进。',.6,6.55,12,.35,17,MUTED)

s=slide('04  验收画面拼图：五类识别与同色干扰','将 acceptance 中的实际图片缩小拼接，集中展示识别结果与复杂背景','来源：acceptance/recording_yolov5_smac_v19/evidence；各帧编号见拼图。','上排和下排第一张覆盖五类图案，其余画面展示蓝色起点、黄色车身和红箱与墙面图案共存。红箱的紫色框来自专用障碍检测器，不是 FIRE 灾害分类。拼图展示本次运行的样例，不等于整段视频完全没有瞬时误框。')
picture(s,AS/'YOLOv5识别与同色干扰拼图.png',.4,1.97,12.5,4.91)

s=slide('创新落地：用同一次验收证明完整流程','感知 → 清障 → 入室定位 → 自主返航','来源：v19/evidence/evaluation.json；模型指标来自当前权重 JSON，两类证据分开引用。','四个创新围绕救援流程协作。v19 同次验收五类全部完成、四个核心标记完成、箱体位移 0.9749 米、返航误差 3.19 厘米。模型独立合成验证精确率 85.62%、召回率 70.81%，不能当成全程视频准确率。没有统一条件前后对照，不给出改进百分比。')
for i,(a,b) in enumerate([('5 / 5','走廊类别'),('4 / 4','核心 AprilTag'),('0.9749 m','箱体实际位移'),('3.19 cm','返航误差')]):
 x=.5+i*3.15;box(s,x,2.25,3,1.8);txt(s,a,x+.18,2.6,2.7,.65,30,TEAL,True);txt(s,b,x+.18,3.38,2.7,.4,17)
txt(s,'独立合成验证：Precision 85.62%  /  Recall 70.81%',.65,4.45,12,.55,24,bold=True)
txt(s,'160 张验证图；IoU ≥ 0.50；置信度 0.45；59 张无目标图中 2 张误报。',.65,5.15,12,.55,18,MUTED)
txt(s,'特色创新：主动观察、受约束接触、局部地图修正、视觉多级核验。',.65,6.15,12,.65,22,GOLD)
r.save(OUT/'AVS作品特色创新点.pptx')
print(OUT/'AVS作品特色创新点.pptx')
